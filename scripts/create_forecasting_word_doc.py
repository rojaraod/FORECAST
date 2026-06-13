"""
Create a Word document for the Supply Chain Emissions Forecasting Framework.

The script converts the canonical Markdown framework into a formatted .docx
document so business and audit teams can review or circulate the methodology in
Microsoft Word.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable, List

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


SOURCE_MARKDOWN = Path("docs/supply_chain_emissions_forecasting_framework.md")
OUTPUT_DOCX = Path("docs/supply_chain_emissions_forecasting_framework.docx")

NAVY = RGBColor(21, 43, 77)
TEAL = RGBColor(0, 128, 128)
GRAY = RGBColor(93, 101, 113)


def clean_inline_markdown(text: str) -> str:
    """Remove simple inline Markdown markers while preserving readable text."""
    cleaned = text.strip()
    cleaned = cleaned.replace("**", "")
    cleaned = cleaned.replace("`", "")
    cleaned = cleaned.replace("<br>", "\n")
    return cleaned


def is_table_separator(line: str) -> bool:
    """Return True when a Markdown line is a table separator row."""
    stripped = line.strip()
    if not stripped.startswith("|") or not stripped.endswith("|"):
        return False
    without_pipes = stripped.replace("|", "").strip()
    return bool(without_pipes) and set(without_pipes) <= {"-", ":", " "}


def is_table_line(line: str) -> bool:
    """Return True for Markdown pipe-table rows."""
    stripped = line.strip()
    return stripped.startswith("|") and stripped.endswith("|") and not is_table_separator(line)


def split_table_row(line: str) -> List[str]:
    """Split a Markdown table row into cleaned cells."""
    return [clean_inline_markdown(cell) for cell in line.strip().strip("|").split("|")]


def add_table(document: Document, table_lines: Iterable[str]) -> None:
    """Add a Markdown table to the Word document."""
    rows = [split_table_row(line) for line in table_lines if is_table_line(line)]
    if not rows:
        return

    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.autofit = True

    for row_idx, row in enumerate(rows):
        for col_idx in range(column_count):
            cell = table.cell(row_idx, col_idx)
            cell.text = row[col_idx] if col_idx < len(row) else ""
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8)
                    if row_idx == 0:
                        run.font.bold = True
                        run.font.color.rgb = NAVY

    document.add_paragraph()


def add_code_block(document: Document, lines: List[str]) -> None:
    """Add a formatted code block."""
    paragraph = document.add_paragraph()
    paragraph.style = "No Spacing"
    for idx, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(8.5)
        if idx < len(lines) - 1:
            run.add_break()


def configure_document(document: Document) -> None:
    """Configure page layout and default styles."""
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.0)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    for style_name, size, color in [
        ("Title", 24, NAVY),
        ("Heading 1", 18, NAVY),
        ("Heading 2", 15, TEAL),
        ("Heading 3", 12, NAVY),
    ]:
        style = styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True


def add_cover_page(document: Document) -> None:
    """Add a simple professional cover page."""
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Supply Chain Emissions Forecasting Framework")
    run.font.name = "Aptos Display"
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = NAVY

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run(
        "End-to-end methodology for supplier-level emissions forecasting, "
        "targets, benchmarks, risk, and dashboards"
    )
    sub_run.font.name = "Aptos"
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = GRAY

    document.add_paragraph()
    for item in [
        "Audience: ESG, Procurement, Sustainability, Data Science, Finance, and Audit teams",
        "Implementation: Python, Excel, and Qlik Sense",
        "Core design: direct historical/current emissions with validation controls",
        "Supplemental benchmarks: historical trend, spend-based emissions, and revenue-based emissions",
    ]:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item).font.name = "Aptos"

    document.add_page_break()


def add_manual_toc(document: Document, markdown_lines: List[str]) -> None:
    """Add a simple manual table of contents from Markdown headings."""
    document.add_heading("Table of Contents", level=1)
    for line in markdown_lines:
        stripped = line.strip()
        if stripped.startswith("## "):
            document.add_paragraph(clean_inline_markdown(stripped[3:]), style="List Bullet")
        elif stripped.startswith("### "):
            paragraph = document.add_paragraph(clean_inline_markdown(stripped[4:]), style="List Bullet 2")
            for run in paragraph.runs:
                run.font.size = Pt(9)
    document.add_page_break()


def markdown_to_docx(markdown_path: Path, output_path: Path) -> None:
    """Convert the framework Markdown file into a formatted Word document."""
    markdown_lines = markdown_path.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)
    add_cover_page(document)
    add_manual_toc(document, markdown_lines)

    idx = 0
    in_code_block = False
    code_lines: List[str] = []

    while idx < len(markdown_lines):
        line = markdown_lines[idx]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code_block:
                add_code_block(document, code_lines)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            idx += 1
            continue

        if in_code_block:
            code_lines.append(line)
            idx += 1
            continue

        if not stripped or stripped == "---":
            idx += 1
            continue

        if is_table_line(line):
            table_lines = []
            while idx < len(markdown_lines) and (
                is_table_line(markdown_lines[idx]) or is_table_separator(markdown_lines[idx])
            ):
                if not is_table_separator(markdown_lines[idx]):
                    table_lines.append(markdown_lines[idx])
                idx += 1
            add_table(document, table_lines)
            continue

        if stripped.startswith("# "):
            document.add_heading(clean_inline_markdown(stripped[2:]), level=1)
        elif stripped.startswith("## "):
            document.add_heading(clean_inline_markdown(stripped[3:]), level=1)
        elif stripped.startswith("### "):
            document.add_heading(clean_inline_markdown(stripped[4:]), level=2)
        elif stripped.startswith("#### "):
            document.add_heading(clean_inline_markdown(stripped[5:]), level=3)
        elif stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(clean_inline_markdown(stripped[2:]))
        elif len(stripped) > 2 and stripped[0].isdigit() and ". " in stripped[:5]:
            paragraph = document.add_paragraph(style="List Number")
            paragraph.add_run(clean_inline_markdown(stripped.split(". ", 1)[1]))
        else:
            paragraph = document.add_paragraph()
            paragraph.add_run(clean_inline_markdown(stripped))

        idx += 1

    if code_lines:
        add_code_block(document, code_lines)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    print(f"Created Word document: {output_path.resolve()}")


if __name__ == "__main__":
    markdown_to_docx(SOURCE_MARKDOWN, OUTPUT_DOCX)
