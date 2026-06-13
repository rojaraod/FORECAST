"""
Create a 20-slide PowerPoint deck for the Supply Chain Emissions Forecasting Model.

The generated presentation is intended for ESG, Procurement, Sustainability,
Data Science, and Audit stakeholders. It summarizes the model objective, source
data, direct emissions validation, forecasting methods, formulas, implementation
approach, dashboard scope, controls, risks, and next steps.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT_PATH = Path("docs/supply_chain_emissions_forecasting_20_slide_deck.pptx")

NAVY = RGBColor(21, 43, 77)
TEAL = RGBColor(0, 128, 128)
GREEN = RGBColor(48, 132, 70)
ORANGE = RGBColor(217, 119, 6)
RED = RGBColor(185, 28, 28)
GRAY = RGBColor(93, 101, 113)
LIGHT_GRAY = RGBColor(242, 245, 248)
WHITE = RGBColor(255, 255, 255)


def set_text_frame(text_frame, paragraphs, font_size=18, color=NAVY, bold_first=False):
    """Populate a text frame with bullet paragraphs."""
    text_frame.clear()
    for idx, paragraph_text in enumerate(paragraphs):
        paragraph = text_frame.paragraphs[0] if idx == 0 else text_frame.add_paragraph()
        paragraph.text = paragraph_text
        paragraph.level = 0
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = color
        paragraph.font.name = "Aptos"
        paragraph.space_after = Pt(7)
        if idx == 0 and bold_first:
            paragraph.font.bold = True


def add_title(slide, title, subtitle=None):
    """Add a consistent slide title and optional subtitle."""
    title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12.25), Inches(0.55))
    title_frame = title_box.text_frame
    title_frame.clear()
    paragraph = title_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.name = "Aptos Display"
    paragraph.font.size = Pt(30)
    paragraph.font.bold = True
    paragraph.font.color.rgb = NAVY

    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.95), Inches(1.25), Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()

    if subtitle:
        subtitle_box = slide.shapes.add_textbox(Inches(0.55), Inches(1.08), Inches(12.0), Inches(0.35))
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.text = subtitle
        subtitle_frame.paragraphs[0].font.name = "Aptos"
        subtitle_frame.paragraphs[0].font.size = Pt(13)
        subtitle_frame.paragraphs[0].font.color.rgb = GRAY


def add_footer(slide, slide_number):
    """Add standard footer text."""
    footer = slide.shapes.add_textbox(Inches(0.55), Inches(7.05), Inches(12.25), Inches(0.25))
    frame = footer.text_frame
    frame.text = f"Supply Chain Emissions Forecasting Model | Slide {slide_number}"
    paragraph = frame.paragraphs[0]
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(9)
    paragraph.font.color.rgb = GRAY
    paragraph.alignment = PP_ALIGN.RIGHT


def add_bullets(slide, x, y, w, h, bullets, font_size=17):
    """Add a bulleted text box."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text_frame(box.text_frame, bullets, font_size=font_size)
    return box


def add_callout(slide, x, y, w, h, title, body, fill_color=LIGHT_GRAY):
    """Add a rounded callout box."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = RGBColor(220, 226, 232)
    frame = shape.text_frame
    frame.margin_left = Inches(0.16)
    frame.margin_right = Inches(0.16)
    frame.margin_top = Inches(0.10)
    frame.margin_bottom = Inches(0.10)
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(15)
    paragraph.font.bold = True
    paragraph.font.color.rgb = NAVY
    body_paragraph = frame.add_paragraph()
    body_paragraph.text = body
    body_paragraph.font.name = "Aptos"
    body_paragraph.font.size = Pt(12)
    body_paragraph.font.color.rgb = GRAY
    return shape


def add_table(slide, x, y, w, h, headers, rows, font_size=10):
    """Add a formatted table."""
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    for col_idx, header in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        paragraph = cell.text_frame.paragraphs[0]
        paragraph.font.color.rgb = WHITE
        paragraph.font.bold = True
        paragraph.font.size = Pt(font_size)
        paragraph.font.name = "Aptos"

    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if row_idx % 2 else LIGHT_GRAY
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.font.color.rgb = NAVY
            paragraph.font.size = Pt(font_size)
            paragraph.font.name = "Aptos"
    return table_shape


def add_process_flow(slide, items, y=2.3):
    """Add a left-to-right process flow."""
    x = 0.55
    width = 1.85
    gap = 0.20
    for idx, (title, body, color) in enumerate(items):
        add_callout(slide, x, y, width, 1.25, title, body, fill_color=color)
        if idx < len(items) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + width - 0.02), Inches(y + 0.43), Inches(gap + 0.12), Inches(0.35))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = GRAY
            arrow.line.fill.background()
        x += width + gap


def add_pillar_cards(slide, cards, y=2.0):
    """Add equally spaced cards for concepts or pillars."""
    x = 0.65
    width = 3.0
    for title, body, color in cards:
        add_callout(slide, x, y, width, 2.0, title, body, fill_color=color)
        x += width + 0.3


def create_presentation() -> None:
    """Build and save the presentation."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    slides = []
    for _ in range(20):
        slides.append(prs.slides.add_slide(blank_layout))

    # 1
    slide = slides[0]
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = NAVY
    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(1.25), Inches(11.7), Inches(1.2))
    title_box.text_frame.text = "Supply Chain Emissions Forecasting Model"
    title_p = title_box.text_frame.paragraphs[0]
    title_p.font.name = "Aptos Display"
    title_p.font.size = Pt(42)
    title_p.font.bold = True
    title_p.font.color.rgb = WHITE
    sub = slide.shapes.add_textbox(Inches(0.78), Inches(2.55), Inches(10.8), Inches(0.8))
    sub.text_frame.text = "Mathematical framework for supplier-level emissions forecasting, targets, risk, and dashboards"
    sub_p = sub.text_frame.paragraphs[0]
    sub_p.font.name = "Aptos"
    sub_p.font.size = Pt(20)
    sub_p.font.color.rgb = RGBColor(207, 235, 235)
    add_callout(slide, 0.8, 4.35, 3.2, 1.0, "Audience", "ESG, Procurement, Sustainability, Data Science, Audit", RGBColor(232, 245, 245))
    add_callout(slide, 4.35, 4.35, 3.2, 1.0, "Output", "Python model, Excel formulas, Qlik Sense dashboard design", RGBColor(232, 245, 245))
    add_callout(slide, 7.9, 4.35, 3.2, 1.0, "Scope", "Supplier target, industry target, and historical trend pathways", RGBColor(232, 245, 245))

    # 2
    slide = slides[1]
    add_title(slide, "Executive Objective", "Forecast supplier emissions from historical data to future target years.")
    add_pillar_cards(
        slide,
        [
            ("Why it matters", "Links supplier emissions to procurement, targets, and decarbonization actions.", RGBColor(232, 245, 245)),
            ("What it produces", "Supplier-year forecasts, glide paths, gaps, risk categories, and confidence scores.", RGBColor(239, 246, 255)),
            ("How it scales", "Same logic can run in Python, Excel, and Qlik Sense.", RGBColor(240, 253, 244)),
        ],
    )
    add_bullets(
        slide,
        0.75,
        4.55,
        11.8,
        1.3,
        [
            "Primary goal: create a transparent, repeatable, and auditable emissions forecasting model.",
            "Business users get target status and risk. Technical users get formulas, flags, and reproducible outputs.",
        ],
        17,
    )
    add_footer(slide, 2)

    # 3
    slide = slides[2]
    add_title(slide, "Business Questions Answered")
    add_table(
        slide,
        0.65,
        1.45,
        12.0,
        4.95,
        ["Stakeholder", "Key question", "Model answer"],
        [
            ["ESG", "Are supplier emissions aligned to targets?", "Forecast vs target and glide path status"],
            ["Procurement", "Which suppliers need engagement?", "Risk category, gap-to-target, top emitters"],
            ["Sustainability", "What reductions are required annually?", "Annual reduction required and carbon budget"],
            ["Data Science", "Which method and assumptions were used?", "Forecast method flag and formula library"],
            ["Audit", "Can calculations be traced?", "Source flags, validation flags, confidence score"],
        ],
        11,
    )
    add_footer(slide, 3)

    # 4
    slide = slides[3]
    add_title(slide, "End-to-End Model Architecture")
    add_process_flow(
        slide,
        [
            ("Source data", "Supplier, activity, emissions, targets, factors", RGBColor(239, 246, 255)),
            ("Emission input", "Use supplied historical and current emissions", RGBColor(232, 245, 245)),
            ("Method select", "Supplier target, industry haircut, historical trend", RGBColor(240, 253, 244)),
            ("Forecast", "Target emissions, glide path, gap, budget", RGBColor(255, 247, 237)),
            ("Outputs", "CSV, Excel, Qlik dashboard, audit flags", RGBColor(254, 242, 242)),
        ],
        y=2.15,
    )
    add_bullets(
        slide,
        0.8,
        4.45,
        11.8,
        1.45,
        [
            "Architecture separates source data, calculations, forecast logic, and dashboard outputs.",
            "Each output record keeps its emissions source, forecast method, quality, and validation flags for traceability.",
        ],
        17,
    )
    add_footer(slide, 4)

    # 5
    slide = slides[4]
    add_title(slide, "Source Data Design")
    add_table(
        slide,
        0.55,
        1.35,
        12.25,
        5.35,
        ["Dataset", "Required fields", "Purpose"],
        [
            ["Supplier Master", "Supplier_ID, Supplier_Name, Industry", "Segmentation, joins, ownership"],
            ["Historical Emissions", "Year, spend, revenue, reported emissions, scopes", "Base emissions and trends"],
            ["Supplier Targets", "Baseline year, target year, reduction percent", "Supplier-specific pathway"],
            ["Industry Haircut", "Industry target year, reduction percent, factors", "Industry pathway assumptions"],
            ["Formula Dictionary", "Calculated parameter, Excel formula", "Excel implementation and audit support"],
        ],
        10,
    )
    add_footer(slide, 5)

    # 6
    slide = slides[5]
    add_title(slide, "Direct Emissions Input Validation")
    add_process_flow(
        slide,
        [
            ("1", "Load supplier emissions", RGBColor(220, 252, 231)),
            ("2", "Validate non-negative values", RGBColor(232, 245, 245)),
            ("3", "Check completeness", RGBColor(239, 246, 255)),
            ("4", "Flag outliers", RGBColor(255, 247, 237)),
            ("5", "Use in forecast", RGBColor(254, 242, 242)),
        ],
        y=1.85,
    )
    add_bullets(
        slide,
        0.75,
        4.25,
        11.8,
        1.55,
        [
            "The 6-level waterfall is not required because all suppliers have historical and current emissions.",
            "Every record receives Emission_Source_Flag, Emissions_Data_Available_Flag, Data_Quality_Flag, Missing_Parameter_Flag, Confidence_Score, and Validation_Flag.",
        ],
        16,
    )
    add_footer(slide, 6)

    # 7
    slide = slides[6]
    add_title(slide, "Data Quality, Confidence, and Validation")
    add_table(
        slide,
        0.65,
        1.35,
        12.0,
        4.8,
        ["Validation item", "Expected result", "Typical confidence", "Audit interpretation"],
        [
            ["Emissions present", "Total_Emission_Value is populated", "0.95", "Supplier-year can be forecast"],
            ["Non-negative", "Emissions are >= 0", "0.95", "Invalid negative values are blocked"],
            ["Reasonable range", "No extreme outlier", "0.90", "Outliers route to review"],
            ["Reporting year present", "Supplier_Year is populated", "0.95", "Trend can be calculated"],
            ["Baseline available", "Baseline year has emissions", "0.90", "Target pathway can be anchored"],
            ["Missing data", "No emissions value", "0.00", "Requires remediation"],
        ],
        10,
    )
    add_footer(slide, 7)

    # 8
    slide = slides[7]
    add_title(slide, "Forecast Method Selection Hierarchy")
    add_callout(slide, 1.0, 1.5, 3.4, 1.5, "1. Supplier Target Pathway", "Use when supplier baseline year, target year, and reduction percent are available.", RGBColor(220, 252, 231))
    add_callout(slide, 4.95, 1.5, 3.4, 1.5, "2. Industry Haircut Pathway", "Use when supplier target is missing but industry target assumptions exist.", RGBColor(232, 245, 245))
    add_callout(slide, 8.9, 1.5, 3.4, 1.5, "3. Historical Trend Pathway", "Use supplier historical emissions CAGR when target pathways are unavailable.", RGBColor(255, 247, 237))
    add_bullets(
        slide,
        0.95,
        4.3,
        11.4,
        1.45,
        [
            "The hierarchy gives preference to supplier commitments, then sector assumptions, then observed historical behavior.",
            "Forecast_Method_Flag makes the chosen pathway transparent in Python, Excel, and Qlik Sense.",
        ],
        17,
    )
    add_footer(slide, 8)

    # 9
    slide = slides[8]
    add_title(slide, "Method 1: Supplier Target Pathway")
    add_bullets(
        slide,
        0.75,
        1.35,
        5.9,
        4.9,
        [
            "Use when supplier-specific target data exists.",
            "Inputs: baseline emissions, baseline year, target year, target reduction percentage.",
            "Target_Emission = Baseline_Emission x (1 - Target_Reduction_Percentage).",
            "Forecast follows the annual glide path to the supplier target.",
            "Best for supplier engagement, contractual targets, and executive accountability.",
        ],
        16,
    )
    add_callout(slide, 7.05, 1.6, 5.2, 1.3, "Business impact", "Aligns forecast with supplier's stated decarbonization commitment.", RGBColor(220, 252, 231))
    add_callout(slide, 7.05, 3.25, 5.2, 1.3, "Advantage", "High relevance and easy to explain to Procurement and ESG teams.", RGBColor(232, 245, 245))
    add_callout(slide, 7.05, 4.9, 5.2, 1.3, "Drawback", "Depends on reliable supplier target evidence and baseline data quality.", RGBColor(255, 247, 237))
    add_footer(slide, 9)

    # 10
    slide = slides[9]
    add_title(slide, "Method 2: Industry Haircut Pathway")
    add_bullets(
        slide,
        0.75,
        1.35,
        5.9,
        4.9,
        [
            "Use when supplier-specific target data is missing but actual supplier emissions are available.",
            "Inputs: industry reduction percentage, industry target year, supplier baseline emissions.",
            "Target_Emission = Baseline_Emission x (1 - Industry_Reduction_Percentage).",
            "Forecast uses a linear glide path from baseline to industry target.",
            "Useful for consistent sector-based assumptions across supplier portfolios.",
        ],
        16,
    )
    add_callout(slide, 7.05, 1.6, 5.2, 1.3, "Business impact", "Gives coverage for suppliers without formal or disclosed targets.", RGBColor(232, 245, 245))
    add_callout(slide, 7.05, 3.25, 5.2, 1.3, "Advantage", "Scalable and simple to maintain by industry.", RGBColor(239, 246, 255))
    add_callout(slide, 7.05, 4.9, 5.2, 1.3, "Drawback", "Less supplier-specific and may hide outperformance or underperformance.", RGBColor(255, 247, 237))
    add_footer(slide, 10)

    # 11
    slide = slides[10]
    add_title(slide, "Method 3: Historical Trend Pathway")
    add_bullets(
        slide,
        0.75,
        1.35,
        5.9,
        4.9,
        [
            "Use when supplier and industry pathway data are unavailable.",
            "Inputs: supplier historical emissions by year and latest emissions.",
            "Historical_Trend = CAGR of supplier emissions history.",
            "Forecast_Emission = Latest_Emission x (1 + Historical_Trend)^n.",
            "Best used as a data-driven trend method with confidence caveats.",
        ],
        16,
    )
    add_callout(slide, 7.05, 1.6, 5.2, 1.3, "Business impact", "Uses actual emissions history to forecast suppliers without target pathways.", RGBColor(255, 247, 237))
    add_callout(slide, 7.05, 3.25, 5.2, 1.3, "Advantage", "Reflects observed behavior and supplier growth patterns.", RGBColor(239, 246, 255))
    add_callout(slide, 7.05, 4.9, 5.2, 1.3, "Drawback", "Sensitive to outliers, reporting changes, and limited history.", RGBColor(254, 242, 242))
    add_footer(slide, 11)

    # 12
    slide = slides[11]
    add_title(slide, "Core Mathematical Formula Library")
    add_table(
        slide,
        0.55,
        1.25,
        12.25,
        5.55,
        ["Parameter", "Formula"],
        [
            ["Total emission", "Use supplied Total_Emission_Value or Reported_Emissions"],
            ["Target emission", "Baseline_Emission x (1 - Reduction_Percentage)"],
            ["Annual reduction", "(Baseline_Emission - Target_Emission) / (Target_Year - Baseline_Year)"],
            ["Glide path", "max(Target_Emission, Baseline_Emission - Annual_Reduction x Years_From_Baseline)"],
            ["Historical forecast", "Latest_Emission x (1 + Historical_CAGR)^Years_After_Latest"],
            ["Gap to target", "Forecast_Emission - Target_Emission"],
            ["Carbon budget remaining", "Target_Carbon_Budget - Cumulative_Emissions"],
        ],
        10,
    )
    add_footer(slide, 12)

    # 13
    slide = slides[12]
    add_title(slide, "Final Forecast Output Dataset")
    add_bullets(
        slide,
        0.75,
        1.35,
        5.8,
        4.9,
        [
            "Supplier identity: ID, name, industry.",
            "Time fields: supplier year, baseline year, target year, forecast year.",
            "Emissions inputs: reported emissions, total emission value, scopes, spend, revenue.",
            "Audit flags: emissions source, data availability, data quality, missing parameters, validation.",
            "Forecast outputs: target, forecast, glide path, gap, budget, confidence, risk, status.",
        ],
        16,
    )
    add_callout(slide, 7.05, 1.55, 5.2, 1.2, "Primary metric", "Forecast_Emission by Supplier_ID and Forecast_Year", RGBColor(232, 245, 245))
    add_callout(slide, 7.05, 3.05, 5.2, 1.2, "Decision metric", "Gap_to_Target and Supplier_Risk_Category", RGBColor(255, 247, 237))
    add_callout(slide, 7.05, 4.55, 5.2, 1.2, "Governance metric", "Confidence_Score and Validation_Flag", RGBColor(239, 246, 255))
    add_footer(slide, 13)

    # 14
    slide = slides[13]
    add_title(slide, "Sample Data and Generated Deliverables")
    add_table(
        slide,
        0.65,
        1.35,
        12.0,
        4.9,
        ["Deliverable", "Description"],
        [
            ["Supplier Master", "100 suppliers across multiple industries"],
            ["Historical Emissions", "1,000 supplier-year records from 2015 to 2024"],
            ["Supplier Targets", "Supplier-specific baseline, target year, and reduction data"],
            ["Industry Haircut Pathway", "Industry reduction assumptions and emission factors"],
            ["Forecast Output", "1,600 supplier-year forecasts from 2025 to 2040"],
            ["Excel Workbook", "All datasets, data dictionary, and formula dictionary"],
        ],
        10,
    )
    add_footer(slide, 14)

    # 15
    slide = slides[14]
    add_title(slide, "Python Implementation Flow")
    add_process_flow(
        slide,
        [
            ("1", "Generate or read data", RGBColor(239, 246, 255)),
            ("2", "Validate emissions inputs", RGBColor(232, 245, 245)),
            ("3", "Create quality and source flags", RGBColor(240, 253, 244)),
            ("4", "Apply forecasts and glide paths", RGBColor(255, 247, 237)),
            ("5", "Export CSV and Excel", RGBColor(254, 242, 242)),
        ],
        y=1.75,
    )
    add_bullets(
        slide,
        0.75,
        4.15,
        11.8,
        1.6,
        [
            "Script: src/supply_chain_emissions_forecast.py",
            "Run: python3 src/supply_chain_emissions_forecast.py --output-dir data/output",
            "Outputs are ready for Excel review and Qlik Sense ingestion.",
        ],
        16,
    )
    add_footer(slide, 15)

    # 16
    slide = slides[15]
    add_title(slide, "Excel Implementation")
    add_bullets(
        slide,
        0.75,
        1.35,
        5.8,
        4.8,
        [
            "Use structured table references for maintainable formulas.",
            "Keep separate tables for source data, targets, factors, and forecast output.",
            "Use flags to make emissions availability and validation status visible to business users.",
            "Protect formula columns and allow source data refresh.",
        ],
        16,
    )
    add_callout(slide, 7.05, 1.45, 5.35, 1.2, "Example", '=IF([@[Reported_Emissions]]<>"",[@[Reported_Emissions]],"Missing Emissions")', RGBColor(239, 246, 255))
    add_callout(slide, 7.05, 3.05, 5.35, 1.2, "Formula dictionary", "Generated as formula_dictionary.csv and included in the Excel workbook.", RGBColor(232, 245, 245))
    add_callout(slide, 7.05, 4.65, 5.35, 1.2, "Audit benefit", "Business users can trace every calculated field to a formula.", RGBColor(240, 253, 244))
    add_footer(slide, 16)

    # 17
    slide = slides[16]
    add_title(slide, "Qlik Sense Dashboard Design")
    add_table(
        slide,
        0.55,
        1.25,
        12.25,
        5.55,
        ["Page", "Key charts", "Business question"],
        [
            ["Landing page", "KPI tiles and navigation", "What is the portfolio view?"],
            ["Executive summary", "Trend, gap, risk, budget", "Are we on track?"],
            ["Supplier trend", "Supplier forecast line, top 20", "Which suppliers drive emissions?"],
            ["Industry benchmark", "Industry emissions and intensity", "Which sectors are highest risk?"],
            ["Gap-to-target", "Gap by supplier and year", "Where is action required?"],
            ["Data quality", "Source, availability, missing fields", "Can we trust the data?"],
            ["Risk heatmap", "Industry x risk category", "Where should Procurement engage?"],
            ["Scenario analysis", "Reduction and growth variables", "What if assumptions change?"],
        ],
        9,
    )
    add_footer(slide, 17)

    # 18
    slide = slides[17]
    add_title(slide, "Controls, Validation, and Auditability")
    add_pillar_cards(
        slide,
        [
            ("Input controls", "Mandatory supplier ID, year, industry; non-negative emissions and activity values.", RGBColor(239, 246, 255)),
            ("Calculation controls", "Direct emissions validation, source flags, validation flags, and method flags.", RGBColor(232, 245, 245)),
            ("Output controls", "Risk category, confidence score, target status, and exception review views.", RGBColor(240, 253, 244)),
        ],
        y=1.55,
    )
    add_bullets(
        slide,
        0.8,
        4.35,
        11.6,
        1.4,
        [
            "Recommended validation: missing emissions, negative values, target year before baseline, low confidence, and outliers.",
            "Each exception should route to supplier data remediation or analyst review.",
        ],
        16,
    )
    add_footer(slide, 18)

    # 19
    slide = slides[18]
    add_title(slide, "Assumptions, Risks, and Limitations")
    add_table(
        slide,
        0.65,
        1.35,
        12.0,
        4.9,
        ["Area", "Assumption or risk", "Mitigation"],
        [
            ["Supplier data", "Reported emissions may be unaudited or inconsistent", "Track source and confidence"],
            ["Factors", "Emission factors may become outdated", "Version factor tables"],
            ["Spend", "Spend growth can reflect inflation, not activity", "Add volume data when available"],
            ["Industry pathways", "Industry averages may not fit individual suppliers", "Replace with supplier targets"],
            ["Historical trends", "Past trend may not predict future action", "Backtest and scenario test"],
        ],
        10,
    )
    add_footer(slide, 19)

    # 20
    slide = slides[19]
    add_title(slide, "Roadmap and Next Steps")
    add_pillar_cards(
        slide,
        [
            ("1. Operationalize", "Load actual supplier data, confirm units, and validate factor sources.", RGBColor(239, 246, 255)),
            ("2. Govern", "Assign data owners, review exceptions, and document factor versions.", RGBColor(232, 245, 245)),
            ("3. Enhance", "Add scenarios, uncertainty ranges, regional factors, and backtesting.", RGBColor(240, 253, 244)),
        ],
        y=1.55,
    )
    add_bullets(
        slide,
        0.85,
        4.35,
        11.5,
        1.5,
        [
            "Recommended decision: approve the model logic as the baseline forecasting method.",
            "Recommended action: pilot with strategic suppliers and prioritize remediation for low-confidence records.",
        ],
        17,
    )
    add_footer(slide, 20)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_PATH)
    print(f"Created presentation: {OUTPUT_PATH.resolve()}")


if __name__ == "__main__":
    create_presentation()
